import sys
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document_5w1h(output_path):
    doc = docx.Document()

    # Set standard margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Styles Palette
    PRIMARY = RGBColor(30, 41, 59)      # Slate 800
    ACCENT = RGBColor(79, 70, 229)      # Indigo 600
    TEXT_DARK = RGBColor(51, 65, 85)    # Slate 700
    TEXT_MUTED = RGBColor(100, 116, 139)# Slate 500
    COLOR_EMERALD = RGBColor(16, 185, 129)

    # Document Header / Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("PROPOSAL & BAHAN PRESENTASI SISTEM ARTALEDGER")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = ACCENT

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Penyajian Komprehensif Arsitektur, Kinerja Finansial & Roadmap ERP Menggunakan Kerangka Analisa 5W + 1H")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = PRIMARY

    # Meta Table (Author, Framework, Version)
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [("Metodologi Analisis", "Kerangka 5W + 1H (What, Why, Who, Where, When, How)"), ("Sistem Utama", "ArtaLedger Enterprise Accounting & ERP")],
        [("Tahun Implementasi", "2026 - 2027"), ("Status Kesiapan", "🟢 Production-Ready (201 Automated Tests Passed 100%)")]
    ]
    for r_idx, row in enumerate(meta_data):
        for c_idx, (k, v) in enumerate(row):
            cell = meta_table.cell(r_idx, c_idx)
            set_cell_background(cell, "F1F5F9" if r_idx % 2 == 0 else "F8FAFC")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            rk = p.add_run(f"{k}: ")
            rk.font.name = "Calibri"
            rk.font.size = Pt(9.5)
            rk.font.bold = True
            rk.font.color.rgb = PRIMARY
            rv = p.add_run(v)
            rv.font.name = "Calibri"
            rv.font.size = Pt(9.5)
            rv.font.color.rgb = TEXT_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Content Sections Helper
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = ACCENT
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = PRIMARY
        return p

    def add_p(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Calibri"
            rb.font.size = Pt(10.5)
            rb.font.bold = True
            rb.font.color.rgb = PRIMARY
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.italic = italic
        r.font.color.rgb = TEXT_DARK
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Calibri"
            rb.font.size = Pt(10)
            rb.font.bold = True
            rb.font.color.rgb = PRIMARY
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = TEXT_DARK
        return p

    # --- RINGKASAN METODE ---
    add_h1("RINGKASAN EKSEKUTIF METODE 5W + 1H")
    add_p("Penyajian materi presentasi sistem ArtaLedger disusun secara terstruktur menggunakan metodologi 5W + 1H guna memberikan pemahaman mendalam, logis, dan terarah bagi seluruh pemangku kepentingan:")
    
    # 5W1H Summary Table
    summary_table = doc.add_table(rows=7, cols=2)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_headers = ["Dimensi 5W + 1H", "Fokus Pembahasan & Intisari Jawaban"]
    for c_idx, h_text in enumerate(s_headers):
        cell = summary_table.cell(0, c_idx)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    s_items = [
        ("1. WHAT (Apa itu ArtaLedger?)", "Sistem Informasi Akuntansi ERP modern berstandar PSAK dengan Core Double-Entry presisi tinggi, modul spesifik rekonsiliasi perbankan, umur piutang/hutang, dan manajemen aset tetap."),
        ("2. WHY (Mengapa Diperlukan?)", "Menjawab risiko human error penjurnalan, ketidakcocokan kertas kerja sistem dengan format Excel auditor, lambatnya rekonsiliasi manual perbankan, dan ancaman manipulasi periode akuntansi."),
        ("3. WHO (Siapa Pemangku Kepentingannya?)", "Jajaran Direksi/C-Level (visibilitas KPI & keputusan strategis), Tim Akuntan (otomasi pembukuan & audit), serta Administrator IT (keamanan, RBAC, dan audit log)."),
        ("4. WHERE (Di Mana Diimplementasikan?)", "Di seluruh unit operasional korporat (Kantor Pusat, Unit Layanan/Klinik, Kantor Cabang) melalui antarmuka web modern reaktif dan arsitektur database terisolasi."),
        ("5. WHEN (Kapan Roadmap & Implementasinya?)", "Sistem fondasi siap produksi saat ini (2026), dilanjutkan roadmap 4 fase strategis sepanjang 2026 - 2027 (Sales/Tax, Inventory/Multi-Bank, Budgeting/API, SaaS/AI)."),
        ("6. HOW (Bagaimana Cara Kerjanya?)", "Bekerja melalui arsitektur ACID Transactions, strict balancing Δ < 0.01, formula dinamis 12 KPI ApexCharts, dan pengujian otomatis terverifikasi 100% (201 Pest Tests).")
    ]
    for r_idx, (dim, val) in enumerate(s_items, start=1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, text_val in enumerate([dim, val]):
            cell = summary_table.cell(r_idx, c_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text_val)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.color.rgb = TEXT_DARK
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = PRIMARY

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- 1. WHAT ---
    add_h1("1. WHAT: Apa Itu ArtaLedger?")
    add_p(
        "ArtaLedger adalah sistem ERP dan akuntansi korporat generasi terbaru yang dirancang untuk mengelola seluruh "
        "siklus pencatatan keuangan korporasi, mulai dari penjurnalan harian, rekonsiliasi perbankan, pembukuan buku besar, "
        "hingga penyajian laporan keuangan berstandar Pernyataan Standar Akuntansi Keuangan (PSAK)."
    )
    add_h2("A. Entitas & Fondasi Perangkat Lunak")
    add_bullet(" Dibangun menggunakan framework Laravel 12 (PHP 8.2+) dengan arsitektur modular Domain-Driven Design (DDD).", "Framework Backend:")
    add_bullet(" Menggunakan Livewire v3 dan Alpine.js yang menghadirkan pengalaman Single Page Application (SPA) tanpa jeda reload halaman.", "Antarmuka Pengguna:")
    add_bullet(" Menggunakan MySQL 8.0 / PostgreSQL dengan kepatuhan penuh ACID (Atomicity, Consistency, Isolation, Durability) dan tipe data DECIMAL(15,2) anti-pembulatan.", "Basis Data Transaksional:")

    add_h2("B. Modul-Modul Fungsional Utama")
    add_bullet(" 12 KPI Baku (EBITDA, NPM, DER, SGA/COGS to Sales, ITO, DSI) dengan grafik tren ApexCharts 12 bulan dan kontrol transaksi bernilai signifikan.", "Dasbor Finansial Eksekutif:")
    add_bullet(" Neraca Lajur 10-Kolom berprinsip Strict Normal Balance Placement, Laba Rugi Komprehensif (OCI PSAK 24), dan Arus Kas Langsung dinamis.", "Paket Pelaporan Keuangan Resmi:")
    add_bullet(" Ekstraksi otomatis rekening koran e-Tax/CMS BRI, matching otomatis 1-ke-1 dan 1-ke-N dengan mutasi kas buku besar.", "Rekonsiliasi Bank Otomatis:")
    add_bullet(" Pelacakan faktur dalam bucket umur (Current s/d >90 Hari) dengan alokasi pembayaran pintar multi-invoice (Smart Settlement M:N).", "Manajemen Umur Piutang & Hutang (Aging AR/AP):")
    add_bullet(" Jadwal penyusutan bulanan otomatis, cetak QR Code/Barcode, dan halaman publik verifikasi aset (/a/{code}).", "Manajemen Aset Tetap & Depresiasi:")
    add_bullet(" Engine impor ribuan baris jurnal Excel dengan pembacaan VLOOKUP real-time dan mekanisme cascade rollback aman.", "Impor Massal Jurnal Excel:")

    # --- 2. WHY ---
    add_h1("2. WHY: Mengapa ArtaLedger Sangat Dibutuhkan?")
    add_p(
        "Sistem akuntansi konvensional dan spreadsheet manual memiliki banyak celah yang menimbulkan kerugian finansial, "
        "keterlambatan pelaporan, dan temuan audit. ArtaLedger hadir secara spesifik untuk mengatasi masalah tersebut:"
    )

    add_h2("A. Latar Belakang Masalah Riil Bisnis")
    add_bullet(" Kesalahan pengetikan nominal atau akun dapat menyebabkan debit dan kredit tidak seimbang, mengacaukan neraca akhir.", "1. Human Error Penjurnalan:")
    add_bullet(" Banyak sistem menyajikan neraca lajur yang susunan saldonya melompat, menyulitkan proses verifikasi auditor eksternal.", "2. Deviasi Kertas Kerja Auditor:")
    add_bullet(" Pencocokan manual mutasi rekening koran bank membutuhkan waktu 3 hingga 5 hari kerja di akhir bulan bagi akuntan.", "3. Beban Rekonsiliasi Bank Lambat:")
    add_bullet(" Aset fisik sering kali hilang atau rusak tanpa tercatat penurunan nilainya di pos beban penyusutan buku besar.", "4. Lepasnya Pengawasan Aset Fisik:")
    add_bullet(" Transaksi pada bulan lalu yang telah diaudit masih dapat diedit secara bebas, memicu potensi kecurangan (fraud).", "5. Kerentanan Manipulasi Data:")

    add_h2("B. Solusi Presisi yang Ditawarkan ArtaLedger")
    add_bullet(" Menolak mutasi jika selisih debit-kredit |Δ| ≥ 0.01, serta memblokir posting langsung ke akun Header/Induk.", "Strict Balancing Rule:")
    add_bullet(" Mengunci saldo akun tetap pada sisi normalnya (Debit/Kredit) meskipun bernilai negatif, menjamin kecocokan 100% dengan lembar kerja Excel auditor.", "Prinsip Excel Parity:")
    add_bullet(" Memangkas waktu rekonsiliasi bulanan dari berhari-hari menjadi hitungan menit secara akurat.", "Otomasi Parser PDF CMS:")
    add_bullet(" Penutupan periode buku resmi (Period Locking) dan penguncian saldo awal neraca (is_locked) yang hanya dapat dibuka oleh Super Admin dengan audit log terekam.", "Audit Trail & Penguncian Periode:")

    # --- 3. WHO ---
    add_h1("3. WHO: Siapa Pemangku Kepentingan Sistem?")
    add_p("ArtaLedger dirancang untuk memberikan dampak positif nyata bagi seluruh tingkatan pemangku kepentingan korporat:")

    add_bullet(" Memperoleh visibilitas menyeluruh terhadap kinerja keuangan real-time melalui 12 KPI eksekutif, grafik tren 12 bulan, dan kontrol langsung terhadap transaksi-transaksi bernilai signifikan tanpa harus menunggu laporan akuntan di akhir bulan.", "1. Jajaran Direksi, Komisaris & C-Level (CEO/CFO):")
    add_bullet(" Mengurangi 80% beban kerja rutin manual, mempercepat rekonsiliasi perbankan, menyajikan kertas kerja neraca lajur yang siap diaudit kapan saja, dan menghilangkan risiko ketidakseimbangan jurnal.", "2. Tim Finansial, Akuntan & Auditor Internal:")
    add_bullet(" Menginput transaksi kas dan mutasi harian secara terstruktur tanpa harus menghafal kode akun akuntansi yang rumit.", "3. Staf Operasional & Kasir Unit Kerja:")
    add_bullet(" Menikmati keunggulan sistem yang stabil, mudah dideploy, terproteksi keamanan AgentShield, bebas dari SQL injection, dan memiliki 201 pengujian otomatis yang menjaga keandalan sistem saat ada pembaruan.", "4. Tim IT & Sistem Informasi Perusahaan:")

    # --- 4. WHERE ---
    add_h1("4. WHERE: Di Mana ArtaLedger Diimplementasikan & Beroperasi?")
    add_p("ArtaLedger memiliki fleksibilitas deployment tinggi untuk mendukung operasional bisnis di berbagai skala geografis:")

    add_h2("A. Lingkup Operasional Organisasi")
    add_bullet(" Pengelolaan konsolidasi laporan keuangan, kontrol bagan akun (COA terpusat), dan persetujuan kebijakan saldo awal.", "Kantor Pusat (Head Office):")
    add_bullet(" Pencatatan transaksi pendapatan dan beban per departemen atau cabang (misal: Unit RS, Klinik Utama, Kantor Perwakilan) secara mandiri.", "Unit Usaha & Kantor Cabang:")
    add_bullet(" Tim auditor dapat mengakses neraca lajur dan buku besar secara read-only secara remote melalui jalur aman.", "Auditor Eksternal & Stakeholder Remote:")

    add_h2("B. Arsitektur Lingkungan Sistem")
    add_bullet(" Dapat dideploy pada server on-premise perusahaan maupun Virtual Private Server (VPS Cloud seperti AWS, DigitalOcean, Alibaba Cloud).", "Infrastruktur Server:")
    add_bullet(" Responsif sempurna diakses melalui desktop, laptop, tablet, hingga smartphone melalui browser modern.", "Akses Antarmuka:")
    add_bullet(" Dilengkapi fitur tunnel aman untuk pengujian eksternal dan publikasi aset fisik barcode tanpa membuka akses database.", "Jalur Tunneling Aman:")

    # --- 5. WHEN ---
    add_h1("5. WHEN: Kapan Jadwal Implementasi & Roadmap Pengembangan?")
    add_p("Sistem telah mencapai fase produksi yang stabil dan siap diekspansi secara bertahap sepanjang 2026 - 2027:")

    # Roadmap Table (1 header + 5 data rows = 6 rows)
    roadmap_table = doc.add_table(rows=6, cols=3)
    roadmap_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    r_headers = ["Fase & Periode", "Fokus Utama Pengembangan", "Output & Deliverables"]
    for c_idx, h_text in enumerate(r_headers):
        cell = roadmap_table.cell(0, c_idx)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    r_items = [
        ("Current State (2026)", "Fondasi Akuntansi Inti (Core GL)", "Double-Entry Engine, Neraca Lajur, P&L PSAK 24, Cash Flow, Rekonsiliasi BRI, Aging AR/AP, Aset Tetap, dan Dasbor KPI (🟢 Ready)."),
        ("Fase 1 (Q4 2026)", "Siklus Bisnis Operasional & Pajak", "Modul Penjualan (Invoicing), Pembelian (PO & Bills), serta Tax Engine PPN 11%/12% dan PPh otomatis terhubung e-Faktur."),
        ("Fase 2 (Q1 2027)", "Manajemen Persediaan & Multi-Bank", "Manajemen Stok Multi-Gudang (FIFO & Moving Average), Stock Opname terpadu, dan Parser Bank BCA, Mandiri, BNI, serta MT940."),
        ("Fase 3 (Q2 2027)", "Kontrol Anggaran & RESTful API", "Modul Budgeting vs Actual dengan Hard Overbudget Guard, Alur Persetujuan Bertingkat, dan OpenAPI v3 untuk POS & E-Commerce."),
        ("Fase 4 (Q3 2027)", "Cloud SaaS & Smart AI Engine", "Arsitektur Multi-Tenancy database terisolasi, Smart OCR Invoice Reader Vision AI, dan AI Accounting Anomaly Detection.")
    ]
    for r_idx, (p_title, p_focus, p_out) in enumerate(r_items, start=1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val_text in enumerate([p_title, p_focus, p_out]):
            cell = roadmap_table.cell(r_idx, c_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val_text)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.color.rgb = TEXT_DARK
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = PRIMARY

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- 6. HOW ---
    add_h1("6. HOW: Bagaimana ArtaLedger Bekerja & Menjamin Mutu?")
    add_p("Keunggulan ArtaLedger bertumpu pada mekanisme kerja yang presisi, pengamanan berlapis, dan metodologi jaminan kualitas perangkat lunak:")

    add_h2("A. Cara Kerja Mekanisme Pembukuan Presisi (Double-Entry Engine)")
    add_bullet(" Setiap mutasi jurnal dibungkus dalam blok transaksi database. Jika terjadi kegagalan sistem di tengah proses, seluruh data dibatalkan secara atomik (Rollback).", "1. ACID Database Transactions:")
    add_bullet(" Sistem memvalidasi bahwa total debit sama dengan total kredit dengan batas toleransi selisih |Δ| < 0.01 sebelum transaksi diizinkan tersimpan sebagai Posting.", "2. Validasi Keseimbangan Real-Time:")
    add_bullet(" Akun berstatus Header (induk klasifikasi) secara otomatis diblokir dari transaksi posting langsung, menjaga konsistensi pohon bagan akun.", "3. Perlindungan Akun Header:")

    add_h2("B. Cara Kerja Mesin Analitik & Pelaporan")
    add_bullet(" Menggunakan service khusus yang mengevaluasi formula dinamis (EBITDA, Rasio Solvabilitas, Likuiditas) dengan sistem memoization agar server tetap ringan.", "1. Dashboard Metric Service:")
    add_bullet(" Data dipetakan otomatis ke kolom saldo normal masing-masing akun, sehingga saldo kontra-akun tetap tersaji pada posisi yang semestinya tanpa merusak format kertas kerja.", "2. Algoritma Strict Normal Balance:")

    add_h2("C. Jaminan Kualitas & Standar Rekayasa Perangkat Lunak")
    add_bullet(" Telah teruji dengan 201 skenario pengujian unit & fitur Pest PHP dengan 749 asersi validasi dan tingkat kelulusan 100%.", "1. 201 Automated Tests (Pest PHP):")
    add_bullet(" Seluruh kode mematuhi standar PSR-12 melalui pemeriksaan otomatis Laravel Pint (0 style violations).", "2. Standar Clean Code (Laravel Pint):")
    add_bullet(" Setiap penambahan rute, service, model, dan pengujian wajib disinkronkan ke dokumen Directed Acyclic Graph (DAG) di KNOWLEDGE_GRAPH.md.", "3. Sinkronisasi Arsitektur (Knowledge Graph Sync):")

    # --- KESIMPULAN ---
    add_h1("7. KESIMPULAN & NILAI STRATEGIS")
    add_p(
        "Melalui pendekatan 5W + 1H, dapat disimpulkan bahwa ArtaLedger bukan sekadar perangkat lunak pencatat angka, "
        "melainkan fondasi transformasi digital finansial korporat yang kokoh, transparan, dan siap berkembang menjadi ERP skala penuh. "
        "Investasi pada sistem ini memberikan kepastian hukum, kepatuhan audit, dan percepatan efisiensi operasional bagi seluruh lini bisnis."
    )

    doc.save(output_path)
    print(f"Document 5W1H successfully created at: {output_path}")

if __name__ == "__main__":
    out_dir = r"d:\Belajar Laravel\artaledger\docs"
    docx_file = os.path.join(out_dir, "Bahan_Presentasi_ArtaLedger.docx")
    docx_5w1h_file = os.path.join(out_dir, "Bahan_Presentasi_ArtaLedger_5W1H.docx")
    
    # Save primary 5W1H file
    create_document_5w1h(docx_5w1h_file)
    
    # Try updating the main file if not locked
    try:
        import shutil
        shutil.copyfile(docx_5w1h_file, docx_file)
        print(f"Also successfully updated main file: {docx_file}")
    except PermissionError:
        print(f"Notice: Main file {docx_file} is currently open in Microsoft Word. Saved to {docx_5w1h_file}.")
